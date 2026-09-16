import adsk.core
import traceback
import os
import sys
import asyncio

_THIS_DIR = os.path.dirname(os.path.abspath(__file__))
_LIB_DIR = os.path.join(_THIS_DIR, 'lib')
for p in (_THIS_DIR, _LIB_DIR):
    if p not in sys.path:
        sys.path.append(p)

from AWSServer import MouserClient
from componentSearchEnginev2 import ComponentSearchEngine
import openpyxl
import nodriver as uc


# Global variables to keep handlers/objects alive
_app = None
_ui = None
_handlers = []

CMD_ID = 'libraryAddin_cmdBtn'
CMD_NAME = 'Generate Electronics Library'
CMD_DESCRIPTION = 'Creates an electronics library from an excel file with mouser numbers'
PANEL_ID = 'SolidScriptsAddinsPanel'  
WORKSPACE_ID = 'FusionSolidEnvironment'
ICON_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'resources')
SESSION_FILE = os.path.join(_THIS_DIR, 'session.json')

# Input IDs
INPUT_DESC = 'descriptionText'
INPUT_MODE = 'sourceModeSelector'
INPUT_PATH = 'filePathInput'
INPUT_BROWSE = 'browseButton'
INPUT_PARTNUM = 'partNumberInput'

MODE_EXCEL = 'From Excel'
MODE_SINGLE = 'Single Component' 

async def execute(inputs):
    '''if not os.path.exists(SESSION_FILE):
        result = _ui.messageBox(
            'This add-in needs permission to view your autodesk account to prevent abuse of the mouser api. This information is not stored\n\n'
            'Click OK to open in browser.',
            'Permissions Required',
            adsk.core.MessageBoxButtonTypes.OKCancelButtonType,
            adsk.core.MessageBoxIconTypes.QuestionIconType
        )
        if result != adsk.core.DialogResults.DialogOK:
            return  # user cancelled, don't proceed to login
    mouser = MouserClient()
    mouser.login()

    #intake and process excel file to python list
    mouserNumbers = []
    mode_input = inputs.itemById(INPUT_MODE)
    selected_mode = mode_input.selectedItem.name if mode_input.selectedItem else MODE_EXCEL
    if selected_mode == MODE_SINGLE:
        mouserNumber = inputs.itemById(INPUT_PARTNUM).value.strip()
        mouserNumbers.append(mouserNumber)
    else:
        excel_path = inputs.itemById(INPUT_PATH).value
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = wb.active
        header_row = None
        header_values = None

        for number, row in enumerate(sheet.iter_rows(values_only=True)):
            if any(cell is not None and cell != " " for cell in row):
                header_row = number
                header_values = row
                break

        header_col = None
        if header_row is not None:
            for col_num, header_name in enumerate(header_values, start=1):
                if header_name == "Mouser #":
                    header_col = col_num
                    break
        else:
            if _ui:
                _ui.messageBox('Failed:\n{}'.format("the file appears to be empty"))
            return
        if header_col == None:
            if _ui:
                _ui.messageBox('Failed:\n{}'.format("there is no \"Mouser #\" header in the file"))
            return

        for row in sheet.iter_rows(min_row=header_row + 2, min_col=header_col, max_col=header_col, values_only=True):
            value = row[0]
            if value is not None:
                mouserNumbers.append(value.strip())

    #use api and convert to a list of part numbers
    mfrNumbers = []
    for mouser_number in mouserNumbers:
        mfrNumbers.append(mouser.search_part(mouser_number))
    #for each part number utilize componentSearchEngine.py
    se = await ComponentSearchEngine.create()
    noInfoList = []
    try:
        await se.login()
        for number in mfrNumbers:
            comp = await se.download_component(number)
            if not comp:
                noInfoList.append(number)
    except Exception as e:
        if _ui:
            _ui.messageBox('Failed:\n{}'.format(e))
    finally:
        await asyncio.sleep(5)
        await se.close()'''
    #create components from downloaded files
    
    #remove all downloads
    #finish
# ---------------------------fusion integration setup--------------------------------
class CommandExecuteHandler(adsk.core.CommandEventHandler):
    def __init__(self):
        super().__init__()

    def notify(self, args):
        try:
            inputs = args.command.commandInputs
            mode_input = inputs.itemById(INPUT_MODE)
            selected_mode = mode_input.selectedItem.name if mode_input.selectedItem else MODE_EXCEL

            if selected_mode == MODE_SINGLE:
                partnum_input = inputs.itemById(INPUT_PARTNUM)
                part_number = partnum_input.value.strip() if partnum_input else ''

                if not part_number:
                    _ui.messageBox('Please enter a Mouser part number before continuing.')
                    args.isValidResult = False
                    return

                uc.loop().run_until_complete(execute(inputs))

            else:
                path_input = inputs.itemById(INPUT_PATH)
                file_path = path_input.value.strip() if path_input else ''

                if not file_path:
                    _ui.messageBox('Please select an Excel file before continuing.')
                    args.isValidResult = False
                    return

                if not os.path.isfile(file_path):
                    _ui.messageBox('The selected file could not be found:\n{}'.format(file_path))
                    args.isValidResult = False
                    return

                uc.loop().run_until_complete(execute(inputs))
        except:
            if _ui:
                _ui.messageBox('Failed:\n{}'.format(traceback.format_exc()))


class CommandInputChangedHandler(adsk.core.InputChangedEventHandler):
    def __init__(self):
        super().__init__()

    def notify(self, args):
        try:
            changed_input = args.input
            inputs = args.inputs

            if changed_input.id == INPUT_BROWSE:
                path_input = inputs.itemById(INPUT_PATH)

                file_dlg = _ui.createFileDialog()
                file_dlg.title = 'Select Excel File'
                file_dlg.filter = 'Excel Files (*.xlsx;*.xls)|*.xlsx;*.xls||'
                file_dlg.isMultiSelectEnabled = False

                dlg_result = file_dlg.showOpen()
                if dlg_result == adsk.core.DialogResults.DialogOK:
                    path_input.value = file_dlg.filename

            elif changed_input.id == INPUT_MODE:
                selected_name = changed_input.selectedItem.name if changed_input.selectedItem else MODE_EXCEL

                path_input = inputs.itemById(INPUT_PATH)
                browse_input = inputs.itemById(INPUT_BROWSE)
                partnum_input = inputs.itemById(INPUT_PARTNUM)

                is_excel_mode = (selected_name == MODE_EXCEL)
                path_input.isVisible = is_excel_mode
                browse_input.isVisible = is_excel_mode
                partnum_input.isVisible = not is_excel_mode
        except:
            if _ui:
                _ui.messageBox('Failed:\n{}'.format(traceback.format_exc()))


class CommandCreatedHandler(adsk.core.CommandCreatedEventHandler):
    def __init__(self):
        super().__init__()

    def notify(self, args):
        try:
            cmd = args.command
            cmd.setDialogInitialSize(400, 200)

            inputs = cmd.commandInputs

            inputs.addTextBoxCommandInput(
                INPUT_DESC,
                '',
                'Select the Excel file containing your Mouser part numbers, ' \
                'or enter a single part number. Then click OK to generate ' \
                'the library.',
                3,
                True  # read-only
            )

            mode_input = inputs.addDropDownCommandInput(
                INPUT_MODE, "Source", adsk.core.DropDownStyles.LabeledIconDropDownStyle
            )
            mode_items = mode_input.listItems
            mode_items.add(MODE_EXCEL, True)   # default selected
            mode_items.add(MODE_SINGLE, False)

            inputs.addStringValueInput(INPUT_PATH, 'Excel File Path', '')
            inputs.addBoolValueInput(INPUT_BROWSE, 'Browse...', False, '', False)

            partnum_input = inputs.addStringValueInput(INPUT_PARTNUM, 'Mouser Part #', '')
            partnum_input.isVisible = False  # hidden until Single Component is selected

            on_execute = CommandExecuteHandler()
            cmd.execute.add(on_execute)
            _handlers.append(on_execute)

            on_input_changed = CommandInputChangedHandler()
            cmd.inputChanged.add(on_input_changed)
            _handlers.append(on_input_changed)
        except:
            if _ui:
                _ui.messageBox('Failed:\n{}'.format(traceback.format_exc()))


def run(context):
    global _app, _ui
    try:
        _app = adsk.core.Application.get()
        _ui = _app.userInterface

        cmd_defs = _ui.commandDefinitions
        cmd_def = cmd_defs.itemById(CMD_ID)
        if not cmd_def:
            cmd_def = cmd_defs.addButtonDefinition(
                CMD_ID,
                CMD_NAME,
                CMD_DESCRIPTION,
                ICON_FOLDER
            )

        on_created = CommandCreatedHandler()
        cmd_def.commandCreated.add(on_created)
        _handlers.append(on_created)

        workspace = _ui.workspaces.itemById(WORKSPACE_ID)
        panel = workspace.toolbarPanels.itemById(PANEL_ID)
        control = panel.controls.itemById(CMD_ID)
        if not control:
            control = panel.controls.addCommand(cmd_def)
            control.isPromoted = True
            control.isPromotedByDefault = True

    except:
        if _ui:
            _ui.messageBox('Failed:\n{}'.format(traceback.format_exc()))


def stop(context):
    try:
        workspace = _ui.workspaces.itemById(WORKSPACE_ID)
        panel = workspace.toolbarPanels.itemById(PANEL_ID)
        control = panel.controls.itemById(CMD_ID)
        if control:
            control.deleteMe()

        cmd_def = _ui.commandDefinitions.itemById(CMD_ID)
        if cmd_def:
            cmd_def.deleteMe()
    except:
        if _ui:
            _ui.messageBox('Failed:\n{}'.format(traceback.format_exc()))